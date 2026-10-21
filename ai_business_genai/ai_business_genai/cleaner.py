"""
cleaner.py
==========
Reusable, generic data-cleaning engine for ANY Excel / CSV file.

No Streamlit code lives here, so the engine can be tested and reused anywhere.

Design rules (important for real company data):
  * SAFE fixes are applied automatically (spaces, blank-like text, number/date formats,
    column names, identical duplicate rows, inconsistent upper/lower case).
  * UNCERTAIN things are NEVER silently changed. They are listed in a "Review_Needed"
    table instead (negative amounts, extreme outliers, same ID with different details,
    values that could not be read as numbers/dates).
  * Missing values are NOT filled with 0. By default they stay blank and are reported.
  * Every removed row is kept in a separate sheet, so nothing is lost.
"""
from __future__ import annotations

import io
import re
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime

import numpy as np
import pandas as pd
from pandas.api import types as ptypes

# --------------------------------------------------------------------------- #
# Constants
# --------------------------------------------------------------------------- #
NULL_TOKENS = {
    "", "na", "n/a", "nan", "null", "none", "nil", "-", "--", "?",
    "#n/a", "#na", "missing", "not available",
}

# Columns that act as identifiers / codes (never converted to numbers, never filled)
KEY_RE = re.compile(r"(^|_)(id|code|sku|no|number)($|_)")
CODE_LIKE_RE = re.compile(
    r"(^|_)(id|code|sku|no|number|phone|mobile|contact|pin|pincode|zip|postal|"
    r"aadhaar|pan|account|ifsc|gst|gstin)($|_)"
)
# Columns where a negative value is suspicious (but may be a legit refund)
MEASURE_RE = re.compile(
    r"(sale|amount|price|revenue|qty|quantity|salary|units|age|cost|total|count|income|payment)"
)
NOT_NEGATIVE_CHECK_RE = re.compile(r"(profit|loss|balance|growth|change|diff|margin|return|variance|net)")
DATE_HINT_RE = re.compile(
    r"(date|dob|birth|created|updated|joined|joining|deadline|expiry|timestamp|datetime)"
)
DATE_LIKE_RE = re.compile(
    r"^\s*(\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}"
    r"|\d{1,2}[-\s/][A-Za-z]{3,9}[-\s/,]*\d{2,4}"
    r"|[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{2,4})"
    r"(\s+\d{1,2}:\d{2}(:\d{2})?)?\s*$"
)
LEADING_ZERO_RE = re.compile(r"^0\d+")
CURRENCY_RE = re.compile(r"(?i)(rs\.?|inr|usd|eur|[₹$€£])")

ISSUE_COLUMNS = ["excel_row", "column", "issue", "value", "note"]


# --------------------------------------------------------------------------- #
# Options / result containers
# --------------------------------------------------------------------------- #
@dataclass
class CleaningOptions:
    standardize_columns: bool = True      # "Sales Amount " -> "sales_amount"
    trim_text: bool = True                # remove extra spaces
    nulls_to_nan: bool = True             # "N/A", "null", "-", "" -> blank
    drop_empty_rows: bool = True
    drop_empty_cols: bool = True
    convert_numbers: bool = True          # "₹1,200" -> 1200
    convert_dates: bool = True            # "15-01-2026" -> real date
    dayfirst: bool = True                 # 05-06-2026 means 5 June (Indian/UK style)
    unify_text_case: bool = True          # north / North / NORTH -> North
    drop_exact_duplicates: bool = True    # rows identical in every column
    missing_strategy: str = "leave"       # leave | median | median_unknown


@dataclass
class CleaningResult:
    df: pd.DataFrame
    original_shape: tuple
    actions: list = field(default_factory=list)
    issues: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(columns=ISSUE_COLUMNS))
    removed_rows: pd.DataFrame = field(default_factory=pd.DataFrame)
    profile: pd.DataFrame = field(default_factory=pd.DataFrame)
    missing_cells: int = 0


# --------------------------------------------------------------------------- #
# Reading files
# --------------------------------------------------------------------------- #
def list_sheets(data: bytes, filename: str) -> list[str]:
    if filename.lower().endswith(".csv"):
        return ["CSV"]
    return pd.ExcelFile(io.BytesIO(data), engine="openpyxl").sheet_names


def _detect_header_row(preview: pd.DataFrame) -> int:
    """Skip title rows such as 'Sales Report 2026' above the real header."""
    if preview.empty:
        return 0
    width = int(preview.notna().sum(axis=1).max())
    if width == 0:
        return 0
    for i in range(len(preview)):
        row = preview.iloc[i].dropna()
        if len(row) >= 0.6 * width and all(isinstance(v, str) for v in row):
            return i
    return 0


def read_file(data: bytes, filename: str, sheet: str | None = None):
    """Return (dataframe, header_row_index). header_row_index is 0-based."""
    if filename.lower().endswith(".csv"):
        last_err = None
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
            try:
                df = pd.read_csv(io.BytesIO(data), sep=None, engine="python",
                                 encoding=enc, skip_blank_lines=False, dtype=str)
                return df, 0
            except UnicodeDecodeError as err:
                last_err = err
        raise last_err  # type: ignore[misc]

    target = sheet if sheet not in (None, "CSV") else 0
    preview = pd.read_excel(io.BytesIO(data), sheet_name=target, header=None,
                            nrows=15, engine="openpyxl")
    header_idx = _detect_header_row(preview)
    # dtype=object keeps cells exactly as typed (no lost leading zeros like 0987...)
    df = pd.read_excel(io.BytesIO(data), sheet_name=target, header=header_idx,
                       engine="openpyxl", dtype=object)
    # restore real int/float/datetime columns; text-like cells (e.g. "0987...") stay text
    df = df.infer_objects()
    return df, header_idx


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #
def _norm(name) -> str:
    return re.sub(r"[^0-9a-z]+", "_", str(name).strip().lower()).strip("_")


def _standardize_columns(columns) -> list[str]:
    names, seen = [], {}
    for col in columns:
        base = _norm(col) or "unnamed_column"
        seen[base] = seen.get(base, 0) + 1
        names.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return names


def _is_text(s: pd.Series) -> bool:
    return s.dtype == object or ptypes.is_string_dtype(s)


def _text_cols(df: pd.DataFrame) -> list:
    return [c for c in df.columns if _is_text(df[c])]


def _is_real_number(s: pd.Series) -> bool:
    return ptypes.is_numeric_dtype(s) and not ptypes.is_bool_dtype(s)


def is_code_like(name) -> bool:
    """True for identifier-style columns (order_id, postal_code, phone ...)."""
    n = _norm(name)
    return bool(CODE_LIKE_RE.search(n) or KEY_RE.search(n))


def numeric_measure_columns(df: pd.DataFrame) -> list:
    """Numeric columns that are real measures (sales, profit ...), not IDs / codes."""
    return [c for c in df.columns if _is_real_number(df[c]) and not is_code_like(c)]


def _parse_number(x):
    """'₹1,200' -> 1200.0 ; '(500)' -> -500.0 ; '15%' -> 15.0 ; junk -> NaN"""
    if x is None or isinstance(x, bool):
        return np.nan
    if isinstance(x, (int, float, np.integer, np.floating)):
        return float(x)
    if not isinstance(x, str):
        return np.nan
    t = x.strip()
    negative = False
    if t.startswith("(") and t.endswith(")"):
        negative, t = True, t[1:-1]
    t = CURRENCY_RE.sub("", t)
    t = t.replace(",", "").replace("%", "").replace(" ", "").replace("\xa0", "")
    if len(t) > 1 and t.endswith("-"):
        negative, t = True, t[:-1]
    try:
        value = float(t)
    except ValueError:
        return np.nan
    if not np.isfinite(value):
        return np.nan
    return -abs(value) if negative else value


def _date_likeness(non_null: pd.Series) -> float:
    def ok(v):
        if isinstance(v, (datetime, date, pd.Timestamp)):
            return True
        return isinstance(v, str) and bool(DATE_LIKE_RE.match(v))
    return float(non_null.map(ok).mean())


def _parse_dates(s: pd.Series, dayfirst: bool) -> pd.Series:
    prepared = s.map(lambda v: v.strip() if isinstance(v, str) else v)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return pd.to_datetime(prepared, errors="coerce", dayfirst=dayfirst, format="mixed")


def _fmt(v) -> str:
    if isinstance(v, (pd.Timestamp, datetime)):
        return v.strftime("%d-%m-%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _kind(s: pd.Series) -> str:
    if ptypes.is_bool_dtype(s):
        return "boolean"
    if ptypes.is_numeric_dtype(s):
        return "number"
    if ptypes.is_datetime64_any_dtype(s):
        return "date"
    return "text"


def profile_columns(df: pd.DataFrame) -> pd.DataFrame:
    n = len(df)
    rows = []
    for c in df.columns:
        s = df[c]
        non_null = s.dropna()
        rows.append({
            "column": c,
            "type": _kind(s),
            "missing": int(s.isna().sum()),
            "missing_%": round(float(s.isna().mean()) * 100, 1) if n else 0.0,
            "unique_values": int(non_null.nunique()),
            "example": _fmt(non_null.iloc[0]) if len(non_null) else "",
        })
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------- #
# The cleaning engine
# --------------------------------------------------------------------------- #
def clean_data(raw: pd.DataFrame, options: CleaningOptions | None = None,
               row_offset: int = 2) -> CleaningResult:
    """
    row_offset converts the DataFrame index to the Excel row number:
    excel_row = index + row_offset  (2 when the header is on row 1).
    """
    opts = options or CleaningOptions()
    df = raw.copy()
    actions: list[str] = []
    issues: list[dict] = []
    removed = pd.DataFrame()

    def add_issues(col, issue, values: pd.Series, note, cap=200):
        total = len(values)
        for idx, val in list(values.items())[:cap]:
            issues.append({"excel_row": int(idx) + row_offset, "column": col,
                           "issue": issue, "value": _fmt(val), "note": note})
        if total > cap:
            issues.append({"excel_row": "", "column": col, "issue": issue, "value": "",
                           "note": f"...and {total - cap} more row(s) not listed"})

    # 1. Column names ------------------------------------------------------- #
    if opts.standardize_columns:
        new_names = _standardize_columns(df.columns)
        if list(map(str, df.columns)) != new_names:
            actions.append("Standardised column names (lowercase_with_underscores)")
        df.columns = new_names
    else:
        df.columns = [str(c) for c in df.columns]

    # 2. Trim spaces -------------------------------------------------------- #
    if opts.trim_text:
        changed = 0
        for c in _text_cols(df):
            old = df[c]
            new = old.map(lambda v: re.sub(r"\s+", " ", v).strip() if isinstance(v, str) else v)
            is_str = old.map(lambda v: isinstance(v, str))
            changed += int((is_str & (old != new)).sum())
            df[c] = new
        if changed:
            actions.append(f"Removed extra spaces in {changed} text cell(s)")

    # 3. Blank-like text -> real blanks ------------------------------------- #
    if opts.nulls_to_nan:
        changed = 0
        for c in _text_cols(df):
            mask = df[c].map(lambda v: isinstance(v, str) and v.strip().lower() in NULL_TOKENS)
            if mask.any():
                changed += int(mask.sum())
                df.loc[mask, c] = np.nan
        if changed:
            actions.append(f"Converted {changed} blank-like cell(s) (N/A, null, '-', ...) to real blanks")

    # 4. Empty columns / rows ---------------------------------------------- #
    if opts.drop_empty_cols:
        empty_cols = [c for c in df.columns if df[c].isna().all()]
        if empty_cols:
            df = df.drop(columns=empty_cols)
            actions.append(f"Removed {len(empty_cols)} completely empty column(s): {empty_cols}")
    if opts.drop_empty_rows:
        before = len(df)
        df = df.dropna(how="all")
        if before != len(df):
            actions.append(f"Removed {before - len(df)} completely empty row(s)")

    # 5. Text -> numbers ---------------------------------------------------- #
    if opts.convert_numbers:
        for c in _text_cols(df):
            if CODE_LIKE_RE.search(_norm(c)):
                continue
            s = df[c]
            non_null = s.dropna()
            if non_null.empty:
                continue
            if non_null.map(lambda v: isinstance(v, str) and bool(LEADING_ZERO_RE.match(v.strip()))).any():
                continue  # looks like codes such as 007 -> keep as text
            parsed = non_null.map(_parse_number)
            if parsed.notna().mean() < 0.9:
                continue
            bad = non_null[parsed.isna()]
            has_percent = bool(non_null.map(lambda v: isinstance(v, str) and "%" in v).any())
            df[c] = s.map(_parse_number).astype(float)
            msg = f"Converted '{c}' from text to numbers"
            if has_percent:
                msg += " (% sign removed, value kept as written, e.g. 15% -> 15)"
            if len(bad):
                msg += f"; {len(bad)} unreadable value(s) left blank - see Review_Needed"
                add_issues(c, "Not a valid number", bad, "Original value could not be read as a number; cell left blank")
            actions.append(msg)

    # 6. Text -> dates ------------------------------------------------------ #
    if opts.convert_dates:
        for c in _text_cols(df):
            non_null = df[c].dropna()
            if non_null.empty:
                continue
            like = _date_likeness(non_null)
            hint = bool(DATE_HINT_RE.search(_norm(c)))
            if like < (0.3 if hint else 0.6):
                continue
            try:
                parsed = _parse_dates(df[c], opts.dayfirst)
            except (ValueError, TypeError):
                continue
            ok = parsed.loc[non_null.index].notna()
            if ok.mean() < (0.5 if hint else 0.8):
                continue
            bad = non_null[~ok]
            df[c] = parsed
            order = "day-first (DD-MM-YYYY)" if opts.dayfirst else "month-first (MM-DD-YYYY)"
            msg = f"Converted '{c}' to real dates, read as {order}"
            if len(bad):
                msg += f"; {len(bad)} invalid date(s) left blank - see Review_Needed"
                add_issues(c, "Invalid date", bad, "Original value is not a recognisable date; cell left blank")
            actions.append(msg)

    # 7. Inconsistent upper/lower case in categories ------------------------ #
    if opts.unify_text_case:
        for c in _text_cols(df):
            if CODE_LIKE_RE.search(_norm(c)):
                continue
            non_null = df[c].dropna()
            if non_null.empty or not non_null.map(lambda v: isinstance(v, str)).all():
                continue
            if non_null.nunique() > max(50, 0.05 * len(non_null)):
                continue  # free text, not a category
            counts = non_null.value_counts()
            canonical: dict[str, str] = {}
            for value in counts.index:  # most frequent spelling first
                canonical.setdefault(value.lower(), value)
            mapping = {v: canonical[v.lower()] for v in counts.index if canonical[v.lower()] != v}
            if mapping:
                df[c] = df[c].map(lambda v: mapping.get(v, v) if isinstance(v, str) else v)
                cells = int(sum(counts[v] for v in mapping))
                actions.append(f"Unified inconsistent capitalisation in '{c}' ({len(mapping)} variant(s), {cells} cell(s))")

    # 8. Exact duplicate rows ---------------------------------------------- #
    if opts.drop_exact_duplicates:
        dup_mask = df.duplicated(keep="first")
        if dup_mask.any():
            removed = df[dup_mask].copy()
            removed.insert(0, "excel_row", removed.index.astype(int) + row_offset)
            df = df[~dup_mask]
            actions.append(f"Removed {int(dup_mask.sum())} exact duplicate row(s) (kept in 'Removed_Duplicates' sheet)")

    # 9. Review flags (never changed automatically) ------------------------- #
    for c in df.columns:  # same ID, different details
        if not KEY_RE.search(_norm(c)):
            continue
        s = df[c]
        non_null = s.dropna()
        if len(non_null) < 3 or non_null.nunique() / len(non_null) < 0.8:
            continue  # IDs that repeat naturally (e.g. customer_id) are not errors
        dup = s.duplicated(keep=False) & s.notna()
        if dup.any():
            add_issues(c, "Same ID appears in several rows with different details", s[dup],
                       "Check whether this is a typo, a second entry or a genuine repeat")

    for c in df.columns:
        if not _is_real_number(df[c]):
            continue
        name = _norm(c)
        if MEASURE_RE.search(name) and not NOT_NEGATIVE_CHECK_RE.search(name) and not KEY_RE.search(name):
            neg = df[c] < 0
            if neg.any():
                add_issues(c, "Negative value", df.loc[neg, c],
                           "Could be a refund/return or a data-entry error - not changed")
        non_null = df[c].dropna()
        if len(non_null) >= 20 and non_null.nunique() > 5 and not CODE_LIKE_RE.search(name):
            q1, q3 = non_null.quantile(0.25), non_null.quantile(0.75)
            iqr = q3 - q1
            if iqr > 0:
                out = (df[c] < q1 - 3 * iqr) | (df[c] > q3 + 3 * iqr)
                if out.any():
                    add_issues(c, "Extreme outlier", df.loc[out, c],
                               "Far outside the normal range (3x IQR) - verify, not changed", cap=50)

    # 10. Missing values ---------------------------------------------------- #
    missing_cells = int(df.isna().sum().sum())
    profile = profile_columns(df)
    if missing_cells:
        actions.append(f"{missing_cells} missing cell(s) found across {int((df.isna().sum() > 0).sum())} column(s)")

    if opts.missing_strategy in ("median", "median_unknown"):
        filled = 0
        for c in df.columns:
            n_missing = int(df[c].isna().sum())
            if not n_missing or ptypes.is_datetime64_any_dtype(df[c]) or KEY_RE.search(_norm(c)):
                continue
            if _is_real_number(df[c]):
                if df[c].notna().any():
                    df[c] = df[c].fillna(df[c].median())
                    filled += n_missing
            elif opts.missing_strategy == "median_unknown":
                df[c] = df[c].fillna("Unknown")
                filled += n_missing
        if filled:
            actions.append(f"Filled {filled} missing cell(s) (numbers = column median"
                           f"{', text = Unknown' if opts.missing_strategy == 'median_unknown' else ''}). "
                           "Dates and ID/code columns are never filled")

    if not actions:
        actions.append("Data was already clean - no rule changed anything")

    issues_df = pd.DataFrame(issues, columns=ISSUE_COLUMNS)
    return CleaningResult(
        df=df.reset_index(drop=True),
        original_shape=raw.shape,
        actions=actions,
        issues=issues_df,
        removed_rows=removed.reset_index(drop=True),
        profile=profile,
        missing_cells=missing_cells,
    )


# --------------------------------------------------------------------------- #
# Excel export
# --------------------------------------------------------------------------- #
def _style_sheet(ws):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fill = PatternFill("solid", start_color="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for i, column in enumerate(ws.columns, start=1):
        longest = max((len(str(c.value)) for c in list(column)[:300] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 50)
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def to_excel_bytes(result: CleaningResult) -> bytes:
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl", datetime_format="DD-MM-YYYY") as writer:
        result.df.to_excel(writer, sheet_name="Cleaned_Data", index=False)
        pd.DataFrame({"step": range(1, len(result.actions) + 1), "what_was_done": result.actions}) \
            .to_excel(writer, sheet_name="Cleaning_Log", index=False)
        review = result.issues if len(result.issues) else pd.DataFrame({"message": ["Nothing needs review"]})
        review.to_excel(writer, sheet_name="Review_Needed", index=False)
        if len(result.removed_rows):
            result.removed_rows.to_excel(writer, sheet_name="Removed_Duplicates", index=False)
        result.profile.to_excel(writer, sheet_name="Column_Profile", index=False)
        for ws in writer.book.worksheets:
            _style_sheet(ws)
    return out.getvalue()
